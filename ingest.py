"""Streaming import for local batch files and bounded web uploads."""
from __future__ import annotations

import csv
import hashlib
import json
import uuid
from pathlib import Path
from typing import BinaryIO, Iterator

from core import extract_post_llm, extract_rules, key
from storage import DATA_DIR, add_records, connection, create_source, mark_import, source_by_sha

SUPPORTED = {".csv", ".xlsx", ".txt", ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".docx"}
BATCH_SIZE = 250
MAX_WEB_MB = 500
MAX_WEB_BYTES = MAX_WEB_MB * 1024 * 1024


def copy_to_store(stream: BinaryIO, name: str, limit_bytes: int | None = None) -> tuple[Path, str]:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"Chưa hỗ trợ định dạng {suffix}")
    folder = DATA_DIR / "originals"
    folder.mkdir(parents=True, exist_ok=True)
    output = folder / f"{uuid.uuid4().hex}{suffix}"
    digest = hashlib.sha256()
    written = 0
    try:
        with output.open("wb") as dst:
            while block := stream.read(1024 * 1024):
                written += len(block)
                if limit_bytes and written > limit_bytes:
                    raise ValueError("Vượt giới hạn tải lên qua giao diện; dùng CLI để nạp tệp lớn")
                digest.update(block)
                dst.write(block)
        if not written:
            raise ValueError("Tệp rỗng")
        return output, digest.hexdigest()
    except Exception:
        output.unlink(missing_ok=True)
        raise


HEADERS = {
    "full_name": {"full_name", "ho ten", "ho va ten", "ten liet si", "ten liet sy"},
    "identifier": {"identifier", "ma ho so", "ma liet si", "so ho so"},
    "birth_year": {"birth_year", "nam sinh", "sinh nam"},
    "hometown": {"hometown", "que quan", "nguyen quan"},
    "military_unit": {"military_unit", "don vi", "don vi chien dau"},
    "rank": {"rank", "cap bac", "quan ham"},
    "enlistment_date": {"enlistment_date", "ngay nhap ngu", "nhap ngu"},
    "death_date": {"death_date", "ngay hy sinh", "hy sinh ngay"},
    "death_place": {"death_place", "noi hy sinh"},
    "death_cause": {"death_cause", "ly do hy sinh"},
    "cemetery": {"cemetery", "noi an tang", "nghia trang", "noi chon cat"},
    "grave_location": {"grave_location", "vi tri mo", "so mo"},
    "relic": {"relic", "di vat", "vat nhan dang"},
    "burial_story": {"burial_story", "loi ke", "loi ke noi chon cat"},
}
HEADER_LOOKUP = {key(alias): field for field, names in HEADERS.items() for alias in names}
TEXT_COLUMNS = {"raw_text", "raw text", "noi dung goc", "noi dung", "van ban"}


def spreadsheet_row(headers: list[str], cells: list) -> tuple[str, dict[str, str], list[dict]]:
    raw = {str(h): "" if c is None else str(c).strip() for h, c in zip(headers, cells)}
    fields = {HEADER_LOOKUP[key(header)]: value for header, value in raw.items()
              if key(header) in HEADER_LOOKUP and value}
    narrative = next((value for header, value in raw.items() if key(header) in TEXT_COLUMNS and value), "")
    # Keep the full row as evidence even when the column mapping is incomplete.
    original = json.dumps(raw, ensure_ascii=False) + ("\n" + narrative if narrative else "")
    if narrative:
        fields = {**extract_rules(narrative).facts(), **fields}
    extras = [{"label": header, "value": value} for header, value in raw.items()
              if value and key(header) not in HEADER_LOOKUP and key(header) not in TEXT_COLUMNS]
    return original, fields, extras


def ocr_image(image) -> str:
    try:
        import pytesseract
        return pytesseract.image_to_string(image, lang="vie+eng")
    except (ImportError, OSError) as exc:
        raise RuntimeError("Cần cài Tesseract OCR, dữ liệu ngôn ngữ vie/eng và pytesseract") from exc
    except Exception as exc:
        raise RuntimeError(f"OCR không hoạt động: {exc}") from exc


def spreadsheet_items(headers: list[str], cells: list, mode: str) -> Iterator[tuple]:
    original, mapped, row_extras = spreadsheet_row(headers, cells)
    has_narrative = any(key(header) in TEXT_COLUMNS and str(value or "").strip()
                        for header, value in zip(headers, cells))
    use_llm = mode == "llm" or (mode == "auto" and (has_narrative or not mapped))
    if not use_llm:
        yield original, mapped, row_extras
        return
    parsed = extract_post_llm(original)
    if not parsed.people:
        yield original, {}, row_extras + [x.model_dump() for x in parsed.post_extras]
        return
    for person in parsed.people:
        fields = person.facts.facts()
        if not (fields.get("full_name") or fields.get("identifier")):
            continue
        extras = [x.model_dump() for x in person.extra_fields]
        if person.evidence_excerpt.strip():
            extras.append({"label": "Trích đoạn đối chiếu", "value": person.evidence_excerpt[:1000]})
        if len(parsed.people) == 1:
            extras.extend(row_extras)
        else:
            extras.extend({"label": "Chưa quy cho cá nhân · " + item["label"], "value": item["value"]}
                          for item in row_extras)
        extras.extend({"label": "Chưa quy cho cá nhân · " + x.label, "value": x.value}
                      for x in parsed.post_extras)
        yield original, fields, extras


def iterate_content(path: Path, mode: str = "auto") -> Iterator[tuple]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        import pandas as pd
        try:
            # Excel CSV exports may use semicolons or tabs, not just commas.
            with path.open(encoding="utf-8-sig", newline="") as source:
                sample = source.read(65536)
            try:
                separator = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                separator = ","
            reader = pd.read_csv(path, chunksize=BATCH_SIZE, dtype=str,
                                 keep_default_na=False, encoding="utf-8-sig", sep=separator)
            for chunk in reader:
                headers = [str(h) for h in chunk.columns]
                for cells in chunk.itertuples(index=False, name=None):
                    if any(str(c).strip() for c in cells):
                        yield from spreadsheet_items(headers, list(cells), mode)
        except UnicodeDecodeError as exc:
            raise ValueError("CSV cần mã hóa UTF-8; hãy lưu lại dưới dạng CSV UTF-8") from exc
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                row_iter = sheet.iter_rows(values_only=True)
                headers = [str(h or "") for h in next(row_iter, [])]
                for cells in row_iter:
                    if not any(c is not None and str(c).strip() for c in cells):
                        continue
                    yield from spreadsheet_items(headers, list(cells), mode)
        finally:
            workbook.close()
    elif suffix == ".pdf":
        import fitz
        from PIL import Image
        with fitz.open(path) as doc:
            for page_number, page in enumerate(doc, 1):
                text = page.get_text("text").strip()
                if len(text) < 40:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    import io
                    text = ocr_image(Image.open(io.BytesIO(pix.tobytes("png"))))
                if text.strip():
                    yield from _expand(f"[Trang {page_number}]\n{text}", mode)
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        from PIL import Image
        with Image.open(path) as image:
            text = ocr_image(image)
        if text.strip():
            yield from _expand(text, mode)
    elif suffix == ".docx":
        from docx import Document
        text = "\n".join(p.text for p in Document(path).paragraphs if p.text.strip())
        if text.strip():
            yield from _expand(text, mode)
    elif suffix == ".txt":
        with path.open(encoding="utf-8-sig") as file:
            # A long text stays in bounded 1000-line sections.
            block = []
            for line in file:
                block.append(line)
                if len(block) >= 1000:
                    text = "".join(block)
                    yield from _expand(text, mode)
                    block.clear()
            if block:
                text = "".join(block)
                yield from _expand(text, mode)
    else:
        raise ValueError("Tệp không được hỗ trợ")


def _expand(text: str, mode: str) -> Iterator[tuple]:
    if mode == "llm":
        parsed = extract_post_llm(text)
        produced = False
        for person in parsed.people:
            fields = person.facts.facts()
            if not (fields.get("full_name") or fields.get("identifier")):
                continue
            produced = True
            extras = [field.model_dump() for field in person.extra_fields]
            extras.extend({"label": "Chưa quy cho cá nhân · " + field.label, "value": field.value}
                          for field in parsed.post_extras)
            if person.evidence_excerpt.strip():
                extras.append({"label": "Trích đoạn đối chiếu", "value": person.evidence_excerpt[:1000]})
            yield text, fields, extras
        if not produced:
            yield text, {}, [field.model_dump() for field in parsed.post_extras]
        return
    if mode not in ("rules", "auto"):
        raise ValueError("Chế độ trích xuất không hợp lệ")
    # Free-form PDF/image/TXT/DOCX always use the model in auto mode.
    if mode == "auto":
        yield from _expand(text, "llm")
    else:
        yield text, extract_rules(text).facts(), []


def _claim_empty_failed_import(source_id: int, mode: str, user_id: int | None) -> dict:
    """Only the same submitter can restart a failed import with zero saved rows.

    Taking the write lock before checking prevents two retries from both
    starting. A partial source with committed records is never reset here.
    """
    with connection() as db:
        db.execute("BEGIN IMMEDIATE")
        source = db.execute("SELECT import_status,submitted_by,file_path FROM sources WHERE id=?",
                            (source_id,)).fetchone()
        count = db.execute("SELECT COUNT(*) FROM records WHERE source_id=?", (source_id,)).fetchone()[0]
        retry = (source["import_status"] == "partial" and count == 0
                 and source["submitted_by"] == user_id)
        if retry:
            db.execute("UPDATE sources SET import_status='processing',import_error='',extraction_mode=? WHERE id=?",
                       (mode, source_id))
        return {"retried": retry, "import_status": source["import_status"],
                "stored_count": count, "file_path": source["file_path"]}


def import_file(path: Path, *, claimed_type: str = "community", provider_role: str = "Nguồn nhập",
                purpose: str = "Cung cấp thông tin", contact: str = "", user_id: int | None = None,
                mode: str = "auto", web_upload: bool = False,
                original_name: str | None = None) -> dict:
    path = Path(path)
    if mode not in ("auto", "rules", "llm"):
        raise ValueError("Chế độ trích xuất không hợp lệ")
    if not path.is_file() or path.suffix.lower() not in SUPPORTED:
        raise ValueError("Không tìm thấy tệp phù hợp")
    with path.open("rb") as stream:
        saved, digest = copy_to_store(stream, original_name or path.name, MAX_WEB_BYTES if web_upload else None)
    existing = source_by_sha(digest)
    retried = False
    old_path = None
    if existing:
        state = _claim_empty_failed_import(existing, mode, user_id)
        if not state["retried"]:
            saved.unlink(missing_ok=True)
            return {"source_id": existing, "count": 0, "duplicate_file": True,
                    "import_status": state["import_status"], "stored_count": state["stored_count"]}
        source_id = existing
        retried = True
        old_path = Path(state["file_path"]) if state["file_path"] else None
    else:
        source_id = create_source("file", original_name or path.name, claimed_type, provider_role,
                                  purpose, contact, user_id, str(saved), digest, extraction_mode=mode)
        mark_import(source_id, "processing")
    buffer: list[tuple] = []
    count = 0
    try:
        if retried:
            if old_path and old_path.is_file():
                saved.unlink(missing_ok=True)
                saved = old_path
            else:
                with connection() as db:
                    db.execute("UPDATE sources SET file_path=? WHERE id=?", (str(saved), source_id))
        for item in iterate_content(saved, mode):
            buffer.append(item)
            if len(buffer) >= BATCH_SIZE:
                add_records(source_id, buffer)
                count += len(buffer)
                buffer.clear()
        if buffer:
            add_records(source_id, buffer)
            count += len(buffer)
        mark_import(source_id, "complete")
    except Exception as exc:
        mark_import(source_id, "partial", str(exc))
        # Already committed rows stay traceable, never reported as a full import.
        raise RuntimeError(f"Nạp một phần ({count} mục), nguồn #{source_id}: {exc}") from exc
    return {"source_id": source_id, "count": count, "duplicate_file": False,
            "retried": retried, "import_status": "complete"}
